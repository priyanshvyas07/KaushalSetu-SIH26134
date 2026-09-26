"""
Document Processing Service for Resumes and Syllabi.
Extracts clean plain text from PDF (via PyMuPDF/fitz) and Word documents (.docx).
"""

import base64
import io
import re
from typing import Optional


class DocumentProcessor:
    @staticmethod
    def extract_text_from_pdf_bytes(pdf_bytes: bytes) -> str:
        """Extract plain text from PDF byte stream using PyMuPDF (fitz) or fallback."""
        try:
            import fitz  # PyMuPDF
            doc = fitz.open(stream=pdf_bytes, filetype="pdf")
            text_chunks = []
            for page in doc:
                text_chunks.append(page.get_text())
            doc.close()
            full_text = "\n".join(text_chunks).strip()
            if full_text:
                return full_text
        except ImportError:
            pass
        except Exception as e:
            print(f"[DocumentProcessor] PyMuPDF error: {e}")

        # Fallback heuristic: Extract printable text blocks from raw PDF stream
        try:
            decoded_latin = pdf_bytes.decode('latin-1', errors='ignore')
            # Extract strings between parentheses (TJ / Tj operators) or text streams
            text_matches = re.findall(r'\((.*?)\)', decoded_latin)
            extracted = " ".join([m for m in text_matches if len(m) > 1 and not m.startswith('/')])
            if len(extracted) > 40:
                return extracted
        except Exception:
            pass

        return ""

    @staticmethod
    def extract_text_from_docx_bytes(docx_bytes: bytes) -> str:
        """Extract plain text from DOCX byte stream using python-docx."""
        try:
            import docx
            doc = docx.Document(io.BytesIO(docx_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join([cell.text.strip() for cell in row.cells if cell.text.strip()])
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
        except ImportError:
            pass
        except Exception as e:
            print(f"[DocumentProcessor] python-docx error: {e}")

        return ""

    @classmethod
    def extract_text_from_base64(cls, base64_str: str, file_type: str = "pdf") -> str:
        """Decode base64 encoded document and extract plain text."""
        try:
            # Strip data URI prefix if present (e.g. data:application/pdf;base64,...)
            if "," in base64_str:
                base64_str = base64_str.split(",", 1)[1]

            raw_bytes = base64.b64decode(base64_str)
            if file_type.lower() == "pdf":
                return cls.extract_text_from_pdf_bytes(raw_bytes)
            elif file_type.lower() in ("docx", "doc"):
                return cls.extract_text_from_docx_bytes(raw_bytes)
            else:
                return raw_bytes.decode("utf-8", errors="ignore")
        except Exception as e:
            print(f"[DocumentProcessor] Base64 decoding failed: {e}")
            return ""


doc_processor = DocumentProcessor()
